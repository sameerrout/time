import os
import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.core.database import get_db_connection
from app.core.logger import get_logger
from app.services.settings_service import SettingsService

logger = get_logger("IssueSlipApp.PurchasedMaterialService")


class PurchasedMaterialService:
    @staticmethod
    def generate_receipt_no() -> str:
        """Generate sequential Inward Purchase Receipt Number (e.g. INW-20261020-0001)."""
        date_str = datetime.datetime.now().strftime("%Y%m%d")
        prefix = f"INW-{date_str}-"
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT receipt_no FROM purchased_materials
                    WHERE receipt_no LIKE ?
                    ORDER BY id DESC LIMIT 1;
                """, (f"{prefix}%",))
                row = cursor.fetchone()
                if row:
                    last_num = int(row["receipt_no"].split("-")[-1])
                    return f"{prefix}{last_num + 1:04d}"
        except Exception as exc:
            logger.error(f"Error generating receipt number: {exc}")
        return f"{prefix}0001"

    @classmethod
    def get_material_stock_snapshot(cls, material_code: str) -> dict:
        """
        Fetch material details and calculate:
        - Opening Stock (current stock before new arrival)
        - Issued Stock (total historical issues from issue slips / ledger)
        - Unit Price (last known price)
        """
        code = material_code.strip()
        result = {
            "material_id": None,
            "material_code": code,
            "material_name": "",
            "uom": "PCS",
            "opening_stock": 0.0,
            "issued_stock": 0.0,
            "closing_stock": 0.0,
            "unit_price": 0.0
        }
        if not code:
            return result

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT id, material_code, name, unit, current_stock, unit_price
                    FROM materials
                    WHERE material_code = ? COLLATE NOCASE
                    LIMIT 1;
                """, (code,))
                mat = cursor.fetchone()
                if mat:
                    result["material_id"] = mat["id"]
                    result["material_code"] = mat["material_code"]
                    result["material_name"] = mat["name"]
                    result["uom"] = mat["unit"] or "PCS"
                    result["opening_stock"] = float(mat["current_stock"])
                    result["unit_price"] = float(mat["unit_price"] or 0.0)

                    # Calculate total issued quantity from issue_slip_items
                    cursor.execute("""
                        SELECT COALESCE(SUM(quantity_issued), 0.0) AS total_issued
                        FROM issue_slip_items
                        WHERE material_id = ?;
                    """, (mat["id"],))
                    iss_row = cursor.fetchone()
                    result["issued_stock"] = float(iss_row["total_issued"]) if iss_row else 0.0
                    result["closing_stock"] = result["opening_stock"]

        except Exception as exc:
            logger.error(f"Error fetching stock snapshot for '{code}': {exc}")

        return result

    @classmethod
    def get_all_records(cls, query: str = "") -> list[dict]:
        """Fetch all purchased material inward entries with S no formatting."""
        records = []
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                if query.strip():
                    q = f"%{query.strip()}%"
                    cursor.execute("""
                        SELECT id, receipt_no, material_id, material_code, material_name, uom,
                               opening_stock, received_stock, issued_stock, closing_stock,
                               unit_price, total_price, COALESCE(gst_percent, 18.0) AS gst_percent,
                               COALESCE(gst_amount, 0.0) AS gst_amount, COALESCE(device_name, '') AS device_name,
                               supplier, challan_no, remarks, created_at
                        FROM purchased_materials
                        WHERE (material_code LIKE ? OR material_name LIKE ? OR receipt_no LIKE ? OR supplier LIKE ?)
                        ORDER BY id DESC;
                    """, (q, q, q, q))
                else:
                    cursor.execute("""
                        SELECT id, receipt_no, material_id, material_code, material_name, uom,
                               opening_stock, received_stock, issued_stock, closing_stock,
                               unit_price, total_price, COALESCE(gst_percent, 18.0) AS gst_percent,
                               COALESCE(gst_amount, 0.0) AS gst_amount, COALESCE(device_name, '') AS device_name,
                               supplier, challan_no, remarks, created_at
                        FROM purchased_materials
                        ORDER BY id DESC;
                    """)

                rows = cursor.fetchall()
                for idx, row in enumerate(rows, start=1):
                    rec = dict(row)
                    rec["s_no"] = idx
                    rec["received_qty"] = rec["received_stock"]
                    gst_p = float(rec.get("gst_percent", 18.0))
                    gst_a = float(rec.get("gst_amount", 0.0))
                    if gst_a == 0.0 and gst_p > 0:
                        gst_a = round(float(rec["received_stock"]) * float(rec["unit_price"]) * (gst_p / 100.0), 2)
                        rec["gst_amount"] = gst_a
                    rec["gst_display"] = f"{gst_p:.0f}%"
                    records.append(rec)
        except Exception as exc:
            logger.error(f"Error reading purchased materials: {exc}")
        return records

    @classmethod
    def record_inward_purchase(cls, data: dict) -> tuple[bool, str, Optional[int]]:
        """
        Record incoming material into the plant:
        1. Auto-registers material if not already existing.
        2. Adds received_stock to materials.current_stock.
        3. Updates material unit_price if provided.
        4. Logs transaction in stock_ledger.
        5. Inserts record into purchased_materials table with GST.
        """
        code = str(data.get("material_code", "")).strip()
        name = str(data.get("material_name", "")).strip()
        uom = str(data.get("uom", "PCS")).strip() or "PCS"
        received_qty = float(data.get("received_qty", data.get("received_stock", 0.0)))
        unit_price = float(data.get("unit_price", 0.0))
        gst_percent = float(data.get("gst_percent", data.get("gst", 18.0)))
        supplier = str(data.get("supplier", "")).strip()
        challan_no = str(data.get("challan_no", "")).strip()
        remarks = str(data.get("remarks", "")).strip()
        device_name = str(data.get("device_name", "")).strip()

        if not code:
            return False, "Material Code is required.", None
        if not name:
            return False, "Material Name is required.", None
        if received_qty <= 0:
            return False, "Received Qty must be greater than zero.", None

        receipt_no = data.get("receipt_no") or cls.generate_receipt_no()
        subtotal = round(received_qty * unit_price, 2)
        gst_amount = round(subtotal * (gst_percent / 100.0), 2)
        total_price = round(subtotal + gst_amount, 2)

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()

                # 1. Look up existing material
                cursor.execute("SELECT id, current_stock, unit_price FROM materials WHERE material_code = ? COLLATE NOCASE;", (code,))
                mat = cursor.fetchone()

                if mat:
                    mat_id = mat["id"]
                    opening_stock = float(mat["current_stock"])
                    new_closing_stock = opening_stock + received_qty

                    # Update stock and unit price
                    cursor.execute("""
                        UPDATE materials
                        SET current_stock = ?,
                            unit_price = CASE WHEN ? > 0 THEN ? ELSE unit_price END,
                            unit = COALESCE(NULLIF(?, ''), unit),
                            name = COALESCE(NULLIF(?, ''), name),
                            updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?;
                    """, (new_closing_stock, unit_price, unit_price, uom, name, mat_id))
                else:
                    # Auto-register new material
                    opening_stock = 0.0
                    new_closing_stock = received_qty
                    barcode = f"BAR-{code}"
                    cursor.execute("""
                        INSERT INTO materials (material_code, name, category, unit, barcode, current_stock, min_stock, unit_price)
                        VALUES (?, ?, 'Purchased Goods', ?, ?, ?, 5.0, ?);
                    """, (code, name, uom, barcode, new_closing_stock, unit_price))
                    mat_id = cursor.lastrowid

                # Calculate total issued quantity
                cursor.execute("""
                    SELECT COALESCE(SUM(quantity_issued), 0.0) AS total_issued
                    FROM issue_slip_items
                    WHERE material_id = ?;
                """, (mat_id,))
                iss_row = cursor.fetchone()
                issued_stock = float(iss_row["total_issued"]) if iss_row else 0.0

                # 2. Log in stock_ledger
                cursor.execute("""
                    INSERT INTO stock_ledger (material_id, transaction_type, quantity, balance_after, reference_no, reason)
                    VALUES (?, 'RECEIPT', ?, ?, ?, ?);
                """, (mat_id, received_qty, new_closing_stock, receipt_no, f"Purchased Material Receipt: {supplier or (f'Phone: {device_name}' if device_name else 'Inward')}"))

                # 3. Insert record into purchased_materials
                cursor.execute("""
                    INSERT INTO purchased_materials (
                        receipt_no, material_id, material_code, material_name, uom,
                        opening_stock, received_stock, issued_stock, closing_stock,
                        unit_price, total_price, gst_percent, gst_amount, device_name, supplier, challan_no, remarks
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                """, (
                    receipt_no, mat_id, code, name, uom,
                    opening_stock, received_qty, issued_stock, new_closing_stock,
                    unit_price, total_price, gst_percent, gst_amount, device_name, supplier, challan_no, remarks
                ))
                record_id = cursor.lastrowid

            logger.info(f"Successfully recorded inward purchase '{receipt_no}' for '{code}' (+{received_qty} {uom}).")
            return True, f"Inward material receipt {receipt_no} recorded successfully.", record_id

        except Exception as exc:
            msg = f"Failed to record inward purchase: {exc}"
            logger.error(msg, exc_info=True)
            return False, msg, None

    @classmethod
    def delete_record(cls, record_id: int) -> tuple[bool, str]:
        """Undo / delete an inward purchase record and revert physical inventory."""
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("SELECT id, material_id, material_code, received_stock, receipt_no FROM purchased_materials WHERE id = ?;", (record_id,))
                rec = cursor.fetchone()
                if not rec:
                    return False, "Record not found."

                mat_id = rec["material_id"]
                received_qty = float(rec["received_stock"])
                receipt_no = rec["receipt_no"]

                # Revert stock in materials table
                cursor.execute("""
                    UPDATE materials
                    SET current_stock = MAX(0.0, current_stock - ?),
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?;
                """, (received_qty, mat_id))

                # Log reversal in ledger
                cursor.execute("SELECT current_stock FROM materials WHERE id = ?;", (mat_id,))
                curr_row = cursor.fetchone()
                balance_after = float(curr_row["current_stock"]) if curr_row else 0.0

                cursor.execute("""
                    INSERT INTO stock_ledger (material_id, transaction_type, quantity, balance_after, reference_no, reason)
                    VALUES (?, 'ADJUSTMENT_SUB', ?, ?, ?, 'Reversal of Inward Receipt');
                """, (mat_id, received_qty, balance_after, f"REV-{receipt_no}"))

                # Delete inward record
                cursor.execute("DELETE FROM purchased_materials WHERE id = ?;", (record_id,))

            logger.info(f"Reverted inward purchase #{record_id} ({receipt_no}).")
            return True, f"Inward purchase record {receipt_no} deleted and stock adjusted."

        except Exception as exc:
            msg = f"Failed to delete inward purchase: {exc}"
            logger.error(msg)
            return False, msg

    @classmethod
    def export_excel(cls, records: list[dict], output_path: str) -> bool:
        """Export purchased material register to formatted Excel file."""
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Purchased Materials"

            ws.views.sheetView[0].showGridLines = True
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
            ws.page_setup.paperSize = ws.PAPERSIZE_A4

            # Fonts & Styles
            font_title = Font(name="Arial", size=14, bold=True, color="000000")
            font_sub = Font(name="Arial", size=9, color="475569")
            font_hdr = Font(name="Arial", size=9, bold=True, color="000000")
            font_data = Font(name="Arial", size=9, color="000000")
            font_total = Font(name="Arial", size=9, bold=True, color="000000")

            fill_hdr = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

            thin_line = Side(style="thin", color="000000")
            border_all = Border(left=thin_line, right=thin_line, top=thin_line, bottom=thin_line)

            align_c = Alignment(horizontal="center", vertical="center")
            align_l = Alignment(horizontal="left", vertical="center")
            align_r = Alignment(horizontal="right", vertical="center")

            company_name = SettingsService.get_setting("company_name", "TIME TECHNOPLAST LTD.")

            # Title Header
            ws.merge_cells("A1:H1")
            ws["A1"] = f"{company_name} - PURCHASED MATERIAL REGISTER"
            ws["A1"].font = font_title
            ws["A1"].alignment = align_c

            ws.merge_cells("A2:H2")
            ws["A2"] = f"Inward Plant Receipts & Stock Tracking | Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
            ws["A2"].font = font_sub
            ws["A2"].alignment = align_c

            ws.row_dimensions[1].height = 24
            ws.row_dimensions[2].height = 18
            ws.row_dimensions[4].height = 26

            # Headers exactly matching user specification
            headers = [
                "S no",
                "material code",
                "material name",
                "uom",
                "recived Qty",
                "unit price",
                "GST",
                "Total price"
            ]

            hdr_row = 4
            for col_idx, h_text in enumerate(headers, start=1):
                cell = ws.cell(row=hdr_row, column=col_idx, value=h_text)
                cell.font = font_hdr
                cell.fill = fill_hdr
                cell.alignment = align_c
                cell.border = border_all

            # Data rows
            curr_row = 5
            tot_received = 0.0
            tot_value = 0.0

            for idx, r in enumerate(records, start=1):
                ws.row_dimensions[curr_row].height = 20

                s_no = idx
                code = r.get("material_code", "")
                name = r.get("material_name", "")
                uom = r.get("uom", "")
                rec_stock = float(r.get("received_qty", r.get("received_stock", 0.0)))
                u_price = float(r.get("unit_price", 0.0))
                gst_p = float(r.get("gst_percent", 18.0))
                t_price = float(r.get("total_price", 0.0))

                tot_received += rec_stock
                tot_value += t_price

                vals = [s_no, code, name, uom, rec_stock, u_price, f"{gst_p:.0f}%", t_price]
                alignments = [align_c, align_c, align_l, align_c, align_r, align_r, align_c, align_r]

                for col_idx, (val, al) in enumerate(zip(vals, alignments), start=1):
                    cell = ws.cell(row=curr_row, column=col_idx, value=val)
                    cell.font = font_data
                    cell.alignment = al
                    cell.border = border_all
                    if col_idx == 5:
                        cell.number_format = "#,##0.0"
                    elif col_idx in (6, 8):
                        cell.number_format = "#,##0.00"

                curr_row += 1

            # Total Summary Row
            ws.row_dimensions[curr_row].height = 22
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
            ws.cell(row=curr_row, column=1, value="TOTALS:").font = font_total
            ws.cell(row=curr_row, column=1).alignment = align_r

            for c in range(1, 5):
                ws.cell(row=curr_row, column=c).border = border_all
                ws.cell(row=curr_row, column=c).fill = fill_total

            c_rec = ws.cell(row=curr_row, column=5, value=tot_received)
            c_rec.font = font_total
            c_rec.alignment = align_r
            c_rec.border = border_all
            c_rec.fill = fill_total
            c_rec.number_format = "#,##0.0"

            ws.cell(row=curr_row, column=6, value="").border = border_all
            ws.cell(row=curr_row, column=6).fill = fill_total

            ws.cell(row=curr_row, column=7, value="").border = border_all
            ws.cell(row=curr_row, column=7).fill = fill_total

            c_val = ws.cell(row=curr_row, column=8, value=tot_value)
            c_val.font = font_total
            c_val.alignment = align_r
            c_val.border = border_all
            c_val.fill = fill_total
            c_val.number_format = "#,##0.00"

            # Column dimensions
            col_widths = {
                1: 8,    # S no
                2: 20,   # material code
                3: 38,   # material name
                4: 10,   # uom
                5: 16,   # recived Qty
                6: 16,   # unit price
                7: 12,   # GST
                8: 18    # Total price
            }
            for col_idx, width in col_widths.items():
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            output_dir = Path(output_path).parent
            output_dir.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            logger.info(f"Exported Purchased Materials Excel to: {output_path}")
            return True

        except Exception as exc:
            logger.error(f"Failed to export Purchased Materials Excel: {exc}", exc_info=True)
            return False

    @classmethod
    def export_pdf(cls, records: list[dict], output_path: str) -> bool:
        """Export purchased material register to printable A4 Landscape PDF."""
        try:
            doc = SimpleDocTemplate(
                output_path,
                pagesize=landscape(A4),
                leftMargin=10 * mm,
                rightMargin=10 * mm,
                topMargin=10 * mm,
                bottomMargin=10 * mm
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle('DocTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=13, alignment=1, spaceAfter=2)
            sub_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontName='Helvetica', fontSize=8, alignment=1, textColor=colors.HexColor('#475569'), spaceAfter=4)
            hdr_style = ParagraphStyle('HdrCell', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, alignment=1)
            cell_style = ParagraphStyle('BodyCell', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9)
            cell_center = ParagraphStyle('BodyCenter', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9, alignment=1)
            cell_right = ParagraphStyle('BodyRight', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9, alignment=2)
            cell_bold_right = ParagraphStyle('BoldRight', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=9, alignment=2)

            company_name = SettingsService.get_setting("company_name", "TIME TECHNOPLAST LTD.")
            elements = [
                Paragraph(f"{company_name} - PURCHASED MATERIAL REGISTER", title_style),
                Paragraph(f"Inward Plant Receipts & Stock Tracking | Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}", sub_style),
                Spacer(1, 3 * mm)
            ]

            headers = [
                "S no",
                "material code",
                "material name",
                "uom",
                "recived\nQty",
                "unit\nprice",
                "GST",
                "Total\nprice"
            ]

            # A4 landscape width is 297mm - 20mm margins = 277mm printable width
            col_widths = [
                14 * mm,  # S no
                32 * mm,  # material code
                92 * mm,  # material name
                18 * mm,  # uom
                30 * mm,  # recived Qty
                30 * mm,  # unit price
                22 * mm,  # GST
                39 * mm   # Total price
            ]

            table_rows = [[Paragraph(h.replace("\n", "<br/>"), hdr_style) for h in headers]]

            tot_received = 0.0
            tot_value = 0.0

            for idx, r in enumerate(records, start=1):
                rec_stock = float(r.get("received_qty", r.get("received_stock", 0.0)))
                u_price = float(r.get("unit_price", 0.0))
                gst_p = float(r.get("gst_percent", 18.0))
                t_price = float(r.get("total_price", 0.0))

                tot_received += rec_stock
                tot_value += t_price

                table_rows.append([
                    Paragraph(str(idx), cell_center),
                    Paragraph(str(r.get("material_code", "")), cell_center),
                    Paragraph(str(r.get("material_name", "")), cell_style),
                    Paragraph(str(r.get("uom", "")), cell_center),
                    Paragraph(f"{rec_stock:.1f}", cell_right),
                    Paragraph(f"{u_price:.2f}", cell_right),
                    Paragraph(f"{gst_p:.0f}%", cell_center),
                    Paragraph(f"{t_price:.2f}", cell_right)
                ])

            # Total row
            table_rows.append([
                Paragraph("<b>TOTALS</b>", cell_center),
                "", "", "",
                Paragraph(f"{tot_received:.1f}", cell_bold_right),
                "", "",
                Paragraph(f"{tot_value:.2f}", cell_bold_right)
            ])

            t = Table(table_rows, colWidths=col_widths, repeatRows=1)
            t.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E2E8F0')),
                ('SPAN', (0, -1), (3, -1)),
            ]))

            elements.append(t)
            doc.build(elements)
            logger.info(f"Exported Purchased Materials PDF to: {output_path}")
            return True

        except Exception as exc:
            logger.error(f"Failed to export Purchased Materials PDF: {exc}", exc_info=True)
            return False
