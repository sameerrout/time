import os
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.config import ICONS_DIR
from app.core.logger import get_logger
from app.services.settings_service import SettingsService

logger = get_logger("IssueSlipApp.IssueSlipExcelService")


class IssueSlipExcelService:
    @classmethod
    def generate_issue_slip_excel(
        cls,
        slip_data: dict,
        items: list[dict],
        output_path: str
    ) -> bool:
        """
        Export a beautifully formatted, ready-to-print Excel workbook (.xlsx)
        matching the TIME TECHNOPLAST LTD. Issue Slip company paper template.
        """
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Issue Slip"

            # Enable grid lines in Excel view
            ws.views.sheetView[0].showGridLines = True

            # Page Setup for A4 Portrait printing
            ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

            # Margin settings (inches)
            ws.page_margins.left = 0.4
            ws.page_margins.right = 0.4
            ws.page_margins.top = 0.5
            ws.page_margins.bottom = 0.5

            # Fonts & Styles
            font_title = Font(name="Arial", size=14, bold=True, color="000000")
            font_addr = Font(name="Arial", size=8.5, color="1E293B")
            font_ref = Font(name="Arial", size=8, bold=True, color="000000")
            font_banner = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            font_meta_lbl = Font(name="Arial", size=9, bold=True, color="000000")
            font_meta_val = Font(name="Arial", size=9, bold=False, color="000000")
            font_tbl_hdr = Font(name="Arial", size=9, bold=True, color="000000")
            font_tbl_data = Font(name="Arial", size=8.5, bold=False, color="000000")
            font_tbl_bold = Font(name="Arial", size=8.5, bold=True, color="000000")
            font_sig = Font(name="Arial", size=8.5, bold=True, color="334155")

            fill_banner = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
            fill_tbl_hdr = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

            thin_line = Side(style="thin", color="000000")
            border_all = Border(left=thin_line, right=thin_line, top=thin_line, bottom=thin_line)
            border_top_only = Border(top=thin_line)

            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
            align_right = Alignment(horizontal="right", vertical="center")

            # Retrieve template settings
            company_name = SettingsService.get_setting("company_name", "TIME TECHNOPLAST LTD.")
            company_address = SettingsService.get_setting(
                "company_address",
                "Survey No. 787/3 & 4, (V) Kishtapur, Near Ravalkol Chaurastha,\n"
                "Shameerpet Road, Medchal Mdl. & Dist. Hyderabad - 501 401 T.S"
            )
            ref_no = SettingsService.get_setting("ref_no", "Ref. No. TTL-HYD/SPIF-09")
            rev_date = SettingsService.get_setting("rev_date", "Rev & Date. 02/01.04.2024")

            # ----------------------------------------------------------------
            # 1. Top Header Row: Company Logo, Title/Address, Ref Box
            # ----------------------------------------------------------------
            ws.row_dimensions[1].height = 28
            ws.row_dimensions[2].height = 18
            ws.row_dimensions[3].height = 18

            # Embed Logo image if exists
            logo_path = ICONS_DIR / "company_logo.png"
            if logo_path.exists():
                try:
                    img = OpenpyxlImage(str(logo_path))
                    img.width = 110
                    img.height = 55
                    ws.add_image(img, "A1")
                except Exception as img_exc:
                    logger.warning(f"Could not embed logo image in Excel: {img_exc}")
                    ws["A1"] = "TIME"
                    ws["A1"].font = Font(name="Arial", size=18, bold=True)
            else:
                ws["A1"] = "TIME"
                ws["A1"].font = Font(name="Arial", size=18, bold=True)

            # Center Company Title & Address
            ws.merge_cells("C1:G1")
            ws["C1"] = company_name
            ws["C1"].font = font_title
            ws["C1"].alignment = align_center

            ws.merge_cells("C2:G3")
            ws["C2"] = company_address
            ws["C2"].font = font_addr
            ws["C2"].alignment = align_center

            # Ref Box (Top Right)
            ws.merge_cells("H1:I1")
            ws["H1"] = ref_no
            ws["H1"].font = font_ref
            ws["H1"].alignment = align_center
            ws["H1"].border = border_all
            ws["I1"].border = border_all

            ws.merge_cells("H2:I2")
            ws["H2"] = rev_date
            ws["H2"].font = font_ref
            ws["H2"].alignment = align_center
            ws["H2"].border = border_all
            ws["I2"].border = border_all

            # ----------------------------------------------------------------
            # 2. Solid Black Banner: ISSUE SLIP
            # ----------------------------------------------------------------
            ws.row_dimensions[4].height = 24
            ws.merge_cells("A4:I4")
            for col in range(1, 10):
                cell = ws.cell(row=4, column=col)
                cell.fill = fill_banner
            ws["A4"] = "ISSUE SLIP"
            ws["A4"].font = font_banner
            ws["A4"].alignment = align_center

            # ----------------------------------------------------------------
            # 3. Two-Column Metadata Block
            # ----------------------------------------------------------------
            dept = slip_data.get("department", "")
            slip_date = slip_data.get("slip_date", "")
            gl_no = slip_data.get("gl_no", "")
            slip_no = slip_data.get("slip_no", "")
            stores_posting_no = slip_data.get("stores_posting_no", "")
            store_posting_date = slip_data.get("store_posting_date", "")
            reservation_no = slip_data.get("reservation_no", "")

            meta_rows = [
                ("Department Name :", dept, "Stores Posting No. :", stores_posting_no),
                ("Date            :", slip_date, "Store Posting Date :", store_posting_date),
                ("G/L No          :", gl_no, "Reservation No.    :", reservation_no),
                ("Issue Slip No.  :", slip_no, "", ""),
            ]

            start_meta_row = 5
            for idx, (lbl_l, val_l, lbl_r, val_r) in enumerate(meta_rows):
                r = start_meta_row + idx
                ws.row_dimensions[r].height = 19

                # Left side
                ws.cell(row=r, column=1, value=lbl_l).font = font_meta_lbl
                ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
                ws.cell(row=r, column=2, value=val_l).font = font_meta_val
                ws.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="center")

                # Right side
                if lbl_r:
                    ws.cell(row=r, column=6, value=lbl_r).font = font_meta_lbl
                    ws.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center")
                    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
                    ws.cell(row=r, column=7, value=val_r).font = font_meta_val
                    ws.cell(row=r, column=7).alignment = Alignment(horizontal="left", vertical="center")

            # ----------------------------------------------------------------
            # 4. Material Table Header (9 Columns matching template)
            # ----------------------------------------------------------------
            tbl_hdr_row = 10
            ws.row_dimensions[tbl_hdr_row].height = 28

            headers = [
                "Sr.\nNo.",
                "Description of\nthe Material",
                "Material Code\nNo.",
                "UOM",
                "Cost Center\nNo.",
                "Required\nQty.",
                "Issued\nQty.",
                "Purpose",
                "Receiver\nName"
            ]

            for c_idx, h_text in enumerate(headers, start=1):
                cell = ws.cell(row=tbl_hdr_row, column=c_idx, value=h_text)
                cell.font = font_tbl_hdr
                cell.fill = fill_tbl_hdr
                cell.alignment = align_center
                cell.border = border_all

            # ----------------------------------------------------------------
            # 5. Material Table Data Rows
            # ----------------------------------------------------------------
            current_row = tbl_hdr_row + 1
            total_issued = 0.0

            for itm_idx, itm in enumerate(items, start=1):
                ws.row_dimensions[current_row].height = 22

                sr_val = itm_idx
                desc_val = str(itm.get("material_name", ""))
                code_val = str(itm.get("material_code", ""))
                uom_val = str(itm.get("unit", ""))
                cc_val = str(itm.get("cost_center_no", slip_data.get("cost_center_no", "")))
                req_val = float(itm.get("required_qty", itm.get("quantity", 0.0)))
                iss_val = float(itm.get("quantity_issued", itm.get("quantity", 0.0)))
                purp_val = str(itm.get("purpose", slip_data.get("purpose", "")))
                recv_val = str(itm.get("receiver_name", slip_data.get("receiver", "")))

                total_issued += iss_val

                row_vals = [sr_val, desc_val, code_val, uom_val, cc_val, req_val, iss_val, purp_val, recv_val]
                aligns = [
                    align_center, align_left, align_center, align_center, align_center,
                    align_right, align_right, align_left, align_left
                ]

                for col_idx, (val, al) in enumerate(zip(row_vals, aligns), start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.font = font_tbl_data
                    cell.alignment = al
                    cell.border = border_all
                    if col_idx in (6, 7):
                        cell.number_format = "#,##0.0"

                current_row += 1

            # Pad with empty rows to preserve paper template appearance (minimum 8 rows)
            min_rows = max(8, len(items))
            while current_row - tbl_hdr_row <= min_rows:
                ws.row_dimensions[current_row].height = 20
                for col_idx in range(1, 10):
                    cell = ws.cell(row=current_row, column=col_idx, value="")
                    cell.border = border_all
                current_row += 1

            # Total Summary Row
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
            lbl_tot = ws.cell(row=current_row, column=1, value="Total Quantity Issued:")
            lbl_tot.font = font_tbl_bold
            lbl_tot.alignment = Alignment(horizontal="right", vertical="center")
            lbl_tot.fill = fill_total

            for c in range(1, 6):
                ws.cell(row=current_row, column=c).border = border_all
                ws.cell(row=current_row, column=c).fill = fill_total

            ws.cell(row=current_row, column=6, value="").border = border_all
            ws.cell(row=current_row, column=6).fill = fill_total

            cell_iss_tot = ws.cell(row=current_row, column=7, value=total_issued)
            cell_iss_tot.font = font_tbl_bold
            cell_iss_tot.alignment = align_right
            cell_iss_tot.border = border_all
            cell_iss_tot.fill = fill_total
            cell_iss_tot.number_format = "#,##0.0"

            ws.cell(row=current_row, column=8, value="").border = border_all
            ws.cell(row=current_row, column=8).fill = fill_total
            ws.cell(row=current_row, column=9, value="").border = border_all
            ws.cell(row=current_row, column=9).fill = fill_total

            current_row += 2

            # ----------------------------------------------------------------
            # 6. Bottom Signatures Section (3 Columns)
            # ----------------------------------------------------------------
            ws.row_dimensions[current_row].height = 24
            ws.row_dimensions[current_row + 1].height = 18

            sig_blocks = [
                (1, 3, "Prepared By / Storekeeper"),
                (4, 6, "Checked By / Quality Dept."),
                (7, 9, "Authorized / Plant Head"),
            ]

            for start_c, end_c, title in sig_blocks:
                # Signature underline
                for c in range(start_c, end_c + 1):
                    ws.cell(row=current_row, column=c).border = border_top_only

                # Label
                ws.merge_cells(start_row=current_row + 1, start_column=start_c, end_row=current_row + 1, end_column=end_c)
                lbl_c = ws.cell(row=current_row + 1, column=start_c, value=title)
                lbl_c.font = font_sig
                lbl_c.alignment = align_center

            # ----------------------------------------------------------------
            # 7. Optimized Column Widths (matching paper proportions)
            # ----------------------------------------------------------------
            col_widths = {
                1: 7,    # Sr. No.
                2: 34,   # Description of the Material
                3: 18,   # Material Code No.
                4: 9,    # UOM
                5: 16,   # Cost Center No.
                6: 14,   # Required Qty.
                7: 14,   # Issued Qty.
                8: 20,   # Purpose
                9: 18,   # Receiver Name
            }
            for col_idx, width in col_widths.items():
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            # Save workbook
            output_dir = Path(output_path).parent
            output_dir.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            logger.info(f"Generated Issue Slip Excel workbook: {output_path}")
            return True

        except Exception as exc:
            logger.error(f"Failed to generate Issue Slip Excel workbook: {exc}", exc_info=True)
            return False
