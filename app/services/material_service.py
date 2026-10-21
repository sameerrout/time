import os
import sqlite3
from typing import Optional
from app.core.database import get_db_connection
from app.core.logger import get_logger

logger = get_logger("IssueSlipApp.MaterialService")


class MaterialService:
    @staticmethod
    def get_by_barcode(barcode: str) -> Optional[dict]:
        """Find an active material by exact barcode (or material code as fallback)."""
        clean_code = barcode.strip()
        if not clean_code:
            return None

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT id, material_code, name, category, unit, barcode, current_stock, min_stock, is_active,
                           COALESCE(unit_price, 0.0) AS unit_price, COALESCE(opening_stock, 0.0) AS opening_stock
                    FROM materials
                    WHERE (barcode = ? OR material_code = ?) AND is_active = 1
                    COLLATE NOCASE
                    LIMIT 1;
                """, (clean_code, clean_code))
                row = cursor.fetchone()
                if row:
                    return dict(row)
        except Exception as exc:
            logger.error(f"Error querying material by barcode '{clean_code}': {exc}")
        return None

    @staticmethod
    def get_by_code(material_code: str) -> Optional[dict]:
        """Find a material by material code."""
        code = material_code.strip()
        if not code:
            return None

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT id, material_code, name, category, unit, barcode, current_stock, min_stock, is_active,
                           COALESCE(unit_price, 0.0) AS unit_price, COALESCE(opening_stock, 0.0) AS opening_stock
                    FROM materials
                    WHERE material_code = ? COLLATE NOCASE
                    LIMIT 1;
                """, (code,))
                row = cursor.fetchone()
                if row:
                    return dict(row)
        except Exception as exc:
            logger.error(f"Error querying material by code '{code}': {exc}")
        return None

    @staticmethod
    def search_materials(query: str = "", active_only: bool = True) -> list[dict]:
        """Search materials by code, name, category, or barcode."""
        q = f"%{query.strip()}%"
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                sql = """
                    SELECT id, material_code, name, category, unit, barcode, current_stock, min_stock, is_active,
                           COALESCE(unit_price, 0.0) AS unit_price, COALESCE(opening_stock, 0.0) AS opening_stock
                    FROM materials
                    WHERE (material_code LIKE ? OR name LIKE ? OR category LIKE ? OR barcode LIKE ?)
                """
                if active_only:
                    sql += " AND is_active = 1"
                sql += " ORDER BY material_code ASC;"

                cursor.execute(sql, (q, q, q, q))
                return [dict(row) for row in cursor.fetchall()]
        except Exception as exc:
            logger.error(f"Error searching materials with query '{query}': {exc}")
            return []

    @staticmethod
    def get_all(active_only: bool = True) -> list[dict]:
        """Retrieve all registered materials."""
        return MaterialService.search_materials("", active_only=active_only)

    @classmethod
    def get_all_summary(cls, query: str = "", active_only: bool = False) -> list[dict]:
        """
        Retrieve material inventory records with full stock columns:
        S no, material code, material name, uom, opening stock, recived stock, issued stock, closing stock, unit price, total price.
        """
        q = f"%{query.strip()}%"
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                sql = """
                    SELECT 
                        m.id,
                        m.material_code,
                        m.name,
                        m.category,
                        m.unit,
                        m.barcode,
                        COALESCE(m.opening_stock, 0.0) AS opening_stock,
                        COALESCE(m.current_stock, 0.0) AS closing_stock,
                        COALESCE(m.unit_price, 0.0) AS unit_price,
                        COALESCE(m.min_stock, 0.0) AS min_stock,
                        m.is_active,
                        (COALESCE(m.received_stock, 0.0) + COALESCE((
                            SELECT SUM(p.received_stock)
                            FROM purchased_materials p
                            WHERE p.material_code = m.material_code COLLATE NOCASE
                        ), (
                            SELECT SUM(sl.quantity)
                            FROM stock_ledger sl
                            WHERE sl.material_id = m.id
                              AND sl.transaction_type IN ('RECEIPT', 'INWARD_PURCHASE')
                              AND sl.reference_no NOT IN ('INITIAL-OPENING', 'OPENING-BALANCE')
                        ), 0.0)) AS received_stock,
                        (COALESCE(m.issued_stock, 0.0) + COALESCE((
                            SELECT SUM(isi.quantity_issued)
                            FROM issue_slip_items isi
                            WHERE isi.material_code = m.material_code COLLATE NOCASE
                        ), (
                            SELECT ABS(SUM(sl.quantity))
                            FROM stock_ledger sl
                            WHERE sl.material_id = m.id
                              AND sl.transaction_type = 'ISSUE'
                        ), 0.0)) AS issued_stock
                    FROM materials m
                    WHERE (m.material_code LIKE ? OR m.name LIKE ? OR m.category LIKE ? OR m.barcode LIKE ?)
                """
                if active_only:
                    sql += " AND m.is_active = 1"
                sql += " ORDER BY m.id ASC;"

                cursor.execute(sql, (q, q, q, q))
                rows = cursor.fetchall()
                results = []
                for row in rows:
                    r = dict(row)
                    closing = float(r.get("closing_stock") or 0.0)
                    unit_p = float(r.get("unit_price") or 0.0)
                    r["total_price"] = round(closing * unit_p, 2)
                    r["opening_stock"] = float(r.get("opening_stock") or 0.0)
                    results.append(r)
                return results
        except Exception as exc:
            logger.error(f"Error fetching materials stock summary: {exc}")
            return []

    @staticmethod
    def add_material(material_code: str, name: str, category: str, unit: str,
                     barcode: str, opening_stock: float = 0.0, min_stock: float = 0.0,
                     unit_price: float = 0.0) -> tuple[bool, str, Optional[int]]:
        """Add a new material with unique code and barcode validation and stock ledger entry."""
        material_code = material_code.strip()
        name = name.strip()
        barcode = barcode.strip()

        if not material_code or not name or not barcode:
            return False, "Material Code, Name, and Barcode are required fields.", None

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                # Check duplicate code
                cursor.execute("SELECT id FROM materials WHERE material_code = ? COLLATE NOCASE;", (material_code,))
                if cursor.fetchone():
                    return False, f"Material Code '{material_code}' is already registered.", None

                # Check duplicate barcode
                cursor.execute("SELECT id FROM materials WHERE barcode = ? COLLATE NOCASE;", (barcode,))
                if cursor.fetchone():
                    return False, f"Barcode '{barcode}' is already assigned to another material.", None

                cursor.execute("""
                    INSERT INTO materials (material_code, name, category, unit, barcode, current_stock, min_stock, opening_stock, unit_price)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                """, (material_code, name, category.strip() or "General", unit.strip() or "PCS",
                      barcode, max(0.0, opening_stock), max(0.0, min_stock), max(0.0, opening_stock), max(0.0, unit_price)))
                mat_id = cursor.lastrowid

                # Record opening stock in ledger if > 0
                if opening_stock > 0:
                    cursor.execute("""
                        INSERT INTO stock_ledger (material_id, transaction_type, quantity, balance_after, reference_no, reason)
                        VALUES (?, 'RECEIPT', ?, ?, 'OPENING-BALANCE', 'Initial stock registration');
                    """, (mat_id, opening_stock, opening_stock))

                logger.info(f"Registered new material: {material_code} ({name}) - Barcode: {barcode}")
                return True, "Material created successfully.", mat_id
        except Exception as exc:
            msg = f"Failed to add material: {exc}"
            logger.error(msg)
            return False, msg, None

    @staticmethod
    def update_material(
        material_id: int,
        name: str,
        category: str,
        unit: str,
        barcode: str,
        min_stock: float,
        is_active: int = 1,
        unit_price: float = 0.0,
        opening_stock: Optional[float] = None
    ) -> tuple[bool, str]:
        """Update existing material details with barcode uniqueness check."""
        name = name.strip()
        barcode = barcode.strip()
        if not name or not barcode:
            return False, "Material Name and Barcode cannot be empty."

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                # Check barcode collision with other items
                cursor.execute("""
                    SELECT id, material_code FROM materials
                    WHERE barcode = ? AND id != ? COLLATE NOCASE;
                """, (barcode, material_id))
                conflict = cursor.fetchone()
                if conflict:
                    return False, f"Barcode '{barcode}' is already in use by material '{conflict['material_code']}'."

                if opening_stock is not None:
                    cursor.execute("""
                        UPDATE materials
                        SET name = ?, category = ?, unit = ?, barcode = ?, min_stock = ?, is_active = ?,
                            unit_price = ?, opening_stock = ?, updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?;
                    """, (name, category.strip() or "General", unit.strip() or "PCS", barcode,
                          max(0.0, min_stock), is_active, max(0.0, unit_price), max(0.0, opening_stock), material_id))
                else:
                    cursor.execute("""
                        UPDATE materials
                        SET name = ?, category = ?, unit = ?, barcode = ?, min_stock = ?, is_active = ?,
                            unit_price = ?, updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?;
                    """, (name, category.strip() or "General", unit.strip() or "PCS", barcode,
                          max(0.0, min_stock), is_active, max(0.0, unit_price), material_id))
                return True, "Material updated successfully."
        except Exception as exc:
            return False, f"Database error updating material: {exc}"

    @staticmethod
    def adjust_stock(
        material_id: int,
        adjustment_type: str,
        qty: float,
        reason: str,
        user_name: str = "Storekeeper"
    ) -> tuple[bool, str]:
        """
        Adjust stock levels with mandatory reason tracking in the immutable stock ledger.
        adjustment_type: 'ADD', 'SUBTRACT', 'SET'
        """
        if qty < 0:
            return False, "Adjustment quantity must be non-negative."
        if not reason.strip():
            return False, "A documented reason is mandatory for inventory stock adjustments."

        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("SELECT current_stock, material_code, name FROM materials WHERE id = ?;", (material_id,))
                mat = cursor.fetchone()
                if not mat:
                    return False, "Material not found."

                current = float(mat["current_stock"])
                if adjustment_type == "ADD":
                    new_stock = current + qty
                    delta = qty
                elif adjustment_type == "SUBTRACT":
                    if qty > current:
                        return False, f"Cannot subtract {qty}. Current stock is only {current}."
                    new_stock = current - qty
                    delta = -qty
                elif adjustment_type == "SET":
                    delta = qty - current
                    new_stock = qty
                else:
                    return False, f"Unknown adjustment type: {adjustment_type}"

                cursor.execute("""
                    UPDATE materials
                    SET current_stock = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?;
                """, (new_stock, material_id))

                cursor.execute("""
                    INSERT INTO stock_ledger (material_id, transaction_type, quantity, balance_after, reference_no, reason)
                    VALUES (?, 'ADJUSTMENT', ?, ?, 'MANUAL-ADJUST', ?);
                """, (material_id, delta, new_stock, f"Adjusted by {user_name}: {reason.strip()}"))

                logger.info(f"Stock adjusted for {mat['material_code']}: {current} -> {new_stock} ({reason})")
                return True, f"Stock updated from {current:.1f} to {new_stock:.1f} successfully."
        except Exception as exc:
            return False, f"Failed to adjust stock: {exc}"

    @staticmethod
    def get_stock_history(material_id: int) -> list[dict]:
        """Retrieve audit history of all movements for this material."""
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT id, transaction_type, quantity, balance_after, reference_no, reason, created_at
                    FROM stock_ledger
                    WHERE material_id = ?
                    ORDER BY id DESC;
                """, (material_id,))
                return [dict(row) for row in cursor.fetchall()]
        except Exception as exc:
            logger.error(f"Error fetching stock history for material {material_id}: {exc}")
            return []

    @classmethod
    def export_to_csv(cls, target_path: str) -> tuple[bool, str]:
        """Export materials stock ledger summary to standard CSV with exact 10 columns."""
        import csv
        try:
            records = cls.get_all_summary("", active_only=False)
            with open(target_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow([
                    "S no", "material code", "material name", "uom",
                    "opening stock", "recived stock", "issued stock",
                    "closing stock", "unit price", "total price"
                ])
                for idx, r in enumerate(records, start=1):
                    writer.writerow([
                        idx,
                        r["material_code"],
                        r["name"],
                        r["unit"],
                        f"{r['opening_stock']:.2f}",
                        f"{r['received_stock']:.2f}",
                        f"{r['issued_stock']:.2f}",
                        f"{r['closing_stock']:.2f}",
                        f"{r['unit_price']:.2f}",
                        f"{r['total_price']:.2f}"
                    ])
            return True, f"Exported {len(records)} materials to CSV."
        except Exception as exc:
            return False, f"CSV Export failed: {exc}"

    @classmethod
    def export_to_excel(cls, target_path: str) -> tuple[bool, str]:
        """Export materials inventory table to formatted Microsoft Excel (.xlsx) file."""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter

            records = cls.get_all_summary("", active_only=False)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Stock Database"
            ws.views.sheetView[0].showGridLines = True

            # Company header
            ws.merge_cells("A1:J1")
            title_cell = ws["A1"]
            title_cell.value = "TIME TECHNOPLAST LTD - STOCK INVENTORY DATABASE"
            title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 30

            headers = [
                "S no", "material code", "material name", "uom",
                "opening stock", "recived stock", "issued stock",
                "closing stock", "unit price", "total price"
            ]

            ws.append([])  # row 2 blank
            ws.append(headers)  # row 3
            ws.row_dimensions[3].height = 24

            thin = Side(border_style="thin", color="CBD5E1")
            cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
            header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

            for col_idx in range(1, 11):
                c = ws.cell(row=3, column=col_idx)
                c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = cell_border

            for idx, r in enumerate(records, start=1):
                row_data = [
                    idx,
                    r["material_code"],
                    r["name"],
                    r["unit"],
                    r["opening_stock"],
                    r["received_stock"],
                    r["issued_stock"],
                    r["closing_stock"],
                    r["unit_price"],
                    r["total_price"]
                ]
                ws.append(row_data)
                curr_row = ws.max_row
                ws.row_dimensions[curr_row].height = 20

                alt_fill = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF",
                                       end_color="F8FAFC" if idx % 2 == 0 else "FFFFFF",
                                       fill_type="solid")
                for c_idx in range(1, 11):
                    cell = ws.cell(row=curr_row, column=c_idx)
                    cell.border = cell_border
                    cell.fill = alt_fill
                    cell.font = Font(name="Calibri", size=10)
                    if c_idx in [1, 2, 4]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif c_idx == 3:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0.00"

            # Auto fit widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row > 1 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(target_path)
            return True, f"Exported {len(records)} materials to Excel at {target_path}"
        except Exception as exc:
            return False, f"Excel Export failed: {exc}"

    @classmethod
    def import_from_csv(cls, source_path: str) -> tuple[int, int, list[str]]:
        """
        Import materials from CSV.
        Expected headers can be:
        1. Exact 10 columns: S no, material code, material name, uom, opening stock, recived stock, issued stock, closing stock, unit price, total price
        2. Legacy: Material Code, Material Name, Category, Unit, Barcode, Current Stock, Min Stock
        Returns (added_count, updated_count, errors)
        """
        import csv
        added, updated = 0, 0
        errors = []

        if not os.path.exists(source_path):
            return 0, 0, ["CSV file does not exist."]

        try:
            with open(source_path, "r", encoding="utf-8-sig", errors="replace") as f:
                reader = csv.DictReader(f)
                
                for line_no, row in enumerate(reader, start=2):
                    row_clean = {k.strip().lower(): v.strip() for k, v in row.items() if k}
                    code = row_clean.get("material code") or row_clean.get("code") or row_clean.get("material_code", "")
                    name = row_clean.get("material name") or row_clean.get("materila name") or row_clean.get("name", "")
                    cat = row_clean.get("category", "General")
                    unit = row_clean.get("uom") or row_clean.get("unit", "PCS")
                    barcode = row_clean.get("barcode", "") or code
                    
                    try:
                        stock = float(row_clean.get("closing stock") or row_clean.get("current stock") or row_clean.get("stock") or 0.0)
                    except ValueError:
                        stock = 0.0
                    try:
                        opening = float(row_clean.get("opening stock") or stock)
                    except ValueError:
                        opening = stock
                    try:
                        unit_p = float(row_clean.get("unit price") or 0.0)
                    except ValueError:
                        unit_p = 0.0
                    try:
                        min_stock = float(row_clean.get("min stock") or row_clean.get("min_stock") or 0.0)
                    except ValueError:
                        min_stock = 0.0

                    if not code or not name:
                        errors.append(f"Line {line_no}: Missing Material Code or Name.")
                        continue

                    existing = cls.get_by_code(code)
                    if existing:
                        # Update existing
                        ok, msg = cls.update_material(
                            existing["id"], name, cat, unit, barcode or existing["barcode"],
                            min_stock, unit_price=unit_p, opening_stock=opening
                        )
                        if ok:
                            updated += 1
                        else:
                            errors.append(f"Line {line_no} ({code}): {msg}")
                    else:
                        # Add new
                        ok, msg, _ = cls.add_material(
                            code, name, cat, unit, barcode,
                            opening_stock=opening, min_stock=min_stock, unit_price=unit_p
                        )
                        if ok:
                            added += 1
                        else:
                            errors.append(f"Line {line_no} ({code}): {msg}")

            return added, updated, errors
        except Exception as exc:
            return added, updated, [f"Failed to read CSV: {exc}"]

    @classmethod
    def import_materials_excel_or_master(cls, file_path: Optional[str] = None) -> tuple[int, int, list[str]]:
        """
        Import material master list from Excel file (columns: Material, Material Description, Base Unit of Measure; Plant 1304 HYDERABAD)
        or fallback to the bundled plant_inventory_master.json (1,734 rows).
        Returns (added_count, updated_count, errors).
        """
        import json
        from pathlib import Path
        from app.core.config import DATA_DIR

        added = 0
        updated = 0
        errors = []

        records = []

        if file_path and Path(file_path).exists():
            path_obj = Path(file_path)
            try:
                if path_obj.suffix.lower() in [".xlsx", ".xls"]:
                    import pandas as pd
                    df = pd.read_excel(file_path)
                    for _, row in df.iterrows():
                        # Map columns
                        code = ""
                        desc = ""
                        uom = "EA"
                        plant = "1304 HYDERABAD"

                        for col in df.columns:
                            c_norm = str(col).strip().lower()
                            val = str(row[col]).strip() if pd.notna(row[col]) else ""
                            if c_norm in ["material", "material code", "code", "mat_code"]:
                                code = val
                            elif c_norm in ["material description", "description", "material name", "name"]:
                                desc = val
                            elif c_norm in ["base unit of measure", "uom", "unit", "base uom"]:
                                uom = val
                            elif c_norm in ["plant"]:
                                plant = val or "1304 HYDERABAD"

                        if code and desc:
                            records.append({
                                "material_code": code,
                                "name": desc,
                                "unit": uom,
                                "plant": plant
                            })
            except Exception as exc:
                errors.append(f"Failed to parse Excel file '{file_path}': {exc}")

        # If no records loaded from Excel, fallback to plant_inventory_master.json
        if not records:
            master_json = DATA_DIR / "plant_inventory_master.json"
            if master_json.exists():
                try:
                    with open(master_json, "r", encoding="utf-8") as f:
                        raw_data = json.load(f)
                        for item in raw_data:
                            records.append({
                                "material_code": item.get("material_code", "").strip(),
                                "name": item.get("material_name", "").strip(),
                                "unit": item.get("uom", "EA").strip(),
                                "opening_stock": float(item.get("opening_stock") or 0.0),
                                "received_stock": float(item.get("received_stock") or 0.0),
                                "issued_stock": float(item.get("issued_stock") or 0.0),
                                "closing_stock": float(item.get("closing_stock") or 0.0),
                                "plant": "1304 HYDERABAD"
                            })
                except Exception as exc:
                    errors.append(f"Failed to parse master JSON: {exc}")

        if not records:
            errors.append("No material records found in Excel or bundled master.")
            return 0, 0, errors

        # Upsert into database
        try:
            with get_db_connection() as conn:
                cursor = conn.cursor()
                for rec in records:
                    code = rec["material_code"]
                    name = rec["name"]
                    unit = rec.get("unit", "EA")
                    plant = rec.get("plant", "1304 HYDERABAD")
                    closing = rec.get("closing_stock", 0.0)
                    opening = rec.get("opening_stock", 0.0)
                    received = rec.get("received_stock", 0.0)
                    issued = rec.get("issued_stock", 0.0)

                    cursor.execute("SELECT id FROM materials WHERE material_code = ? COLLATE NOCASE;", (code,))
                    existing = cursor.fetchone()
                    if existing:
                        cursor.execute("""
                            UPDATE materials
                            SET name = ?, unit = ?, plant = ?
                            WHERE id = ?;
                        """, (name, unit, plant, existing["id"]))
                        updated += 1
                    else:
                        cursor.execute("""
                            INSERT INTO materials (
                                material_code, name, category, unit, barcode,
                                opening_stock, received_stock, issued_stock, current_stock,
                                unit_price, min_stock, is_active, plant
                            ) VALUES (?, ?, 'Plant Materials', ?, ?, ?, ?, ?, ?, 0.0, 0.0, 1, ?);
                        """, (code, name, unit, code, opening, received, issued, closing, plant))
                        added += 1
        except Exception as exc:
            errors.append(f"Database error during material master import: {exc}")

        logger.info(f"Imported material master: {added} added, {updated} updated.")
        return added, updated, errors

