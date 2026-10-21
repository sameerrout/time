import os
import unittest
from pathlib import Path
from PIL import Image
import openpyxl

os.environ["QT_QPA_PLATFORM"] = "offscreen"
from PySide6.QtWidgets import QApplication

from app.services.barcode_service import BarcodeService
from app.ui.views.barcode_generator_view import BarcodeGeneratorView
from app.ui.dialogs.sticker_print_dialog import StickerPrintDialog
from app.core.config import DATA_DIR

app = QApplication.instance() or QApplication([])


class TestPlantQRCatalogAndStickerPrinter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = BarcodeService.get_plant_catalog()

    def test_01_catalog_contains_all_1734_materials(self):
        """Verify all 1,734 plant materials from the 73-page PDF are extracted and loaded."""
        self.assertEqual(len(self.catalog), 1734, "Catalog must contain exactly 1,734 materials without missing any.")

        # Check first and last items
        first_item = self.catalog[0]
        self.assertEqual(first_item["material_code"], "1001NM210010BL421")
        self.assertEqual(first_item["uom"], "EA")
        self.assertIn("INM 210 LTRS", first_item["material_name"])

        last_item = self.catalog[-1]
        self.assertEqual(last_item["material_code"], "4092GENYR000020")
        self.assertEqual(last_item["material_name"], "NYLON SUTALI")
        self.assertEqual(last_item["uom"], "KG")

    def test_02_pure_qr_generation(self):
        """Verify high-contrast QR code image generation for material code."""
        code = "1001NM210010BL421"
        img = BarcodeService.generate_pure_qr_image(code)
        self.assertIsNotNone(img)
        self.assertIsInstance(img, Image.Image)
        self.assertGreater(img.size[0], 50)
        self.assertGreater(img.size[1], 50)

    def test_03_thermal_sticker_label_generation(self):
        """Verify 50x25mm and 50x50mm thermal sticker label image generation."""
        item = self.catalog[0]
        sticker_50x25 = BarcodeService.generate_thermal_sticker_image(
            item["material_code"], item["material_name"], item["uom"], label_size="50x25"
        )
        self.assertIsNotNone(sticker_50x25)
        self.assertGreater(sticker_50x25.size[0], 300)
        self.assertGreater(sticker_50x25.size[1], 150)

        sticker_50x50 = BarcodeService.generate_thermal_sticker_image(
            item["material_code"], item["material_name"], item["uom"], label_size="50x50"
        )
        self.assertIsNotNone(sticker_50x50)
        self.assertEqual(sticker_50x50.size[0], sticker_50x50.size[1])

    def test_04_export_catalog_to_excel_with_qr_codes(self):
        """Verify Excel export matches the user's reference image with embedded QR codes."""
        test_items = self.catalog[:10]
        out_excel = DATA_DIR / "test_catalog_sample.xlsx"
        ok = BarcodeService.export_catalog_to_excel(str(out_excel), items=test_items)
        self.assertTrue(ok)
        self.assertTrue(out_excel.exists())

        wb = openpyxl.load_workbook(str(out_excel))
        ws = wb.active
        # Check column headers match user image
        headers = [ws.cell(row=1, column=i).value for i in range(1, 5)]
        self.assertEqual(headers, ["Material Code", "Material Name", "UOM", "QR Code"])
        # Check first data row
        self.assertEqual(ws.cell(row=2, column=1).value, test_items[0]["material_code"])
        self.assertEqual(ws.cell(row=2, column=2).value, test_items[0]["material_name"])
        self.assertEqual(ws.cell(row=2, column=3).value, test_items[0]["uom"])
        # Check images embedded
        self.assertEqual(len(ws._images), len(test_items))

        wb.close()
        if out_excel.exists():
            out_excel.unlink()

    def test_05_export_sticker_printer_pdf(self):
        """Verify PDF export for thermal continuous roll and A4 sheets."""
        test_items = self.catalog[:3]
        out_pdf_50x25 = DATA_DIR / "test_stickers_50x25.pdf"
        ok = BarcodeService.export_sticker_printer_pdf(str(out_pdf_50x25), test_items, label_size="50x25")
        self.assertTrue(ok)
        self.assertTrue(out_pdf_50x25.exists())
        self.assertGreater(out_pdf_50x25.stat().st_size, 500)
        out_pdf_50x25.unlink()

        out_pdf_a4 = DATA_DIR / "test_stickers_a4.pdf"
        ok2 = BarcodeService.export_sticker_printer_pdf(str(out_pdf_a4), test_items, label_size="a4_sheet")
        self.assertTrue(ok2)
        self.assertTrue(out_pdf_a4.exists())
        out_pdf_a4.unlink()

    def test_06_barcode_generator_view_catalog_mode(self):
        """Verify view loads both modes and table columns match user specifications."""
        view = BarcodeGeneratorView()
        self.assertEqual(view.stacked_widget.count(), 2)

        # Check header button exists in the top right
        self.assertIsNotNone(view.btn_toggle_mode)
        self.assertIsNotNone(view.btn_header_sticker)
        self.assertIsNotNone(view.btn_header_excel)

        # Toggle to Catalog
        view.toggle_view_mode()
        self.assertEqual(view.stacked_widget.currentIndex(), 1)
        self.assertEqual(view.table_catalog.columnCount(), 6)

        # Check table headers match: S no | Material Code | Material Name | UOM | QR Code | Actions
        expected_cols = ["S no", "Material Code", "Material Name", "UOM", "QR Code", "Actions"]
        actual_cols = [view.table_catalog.horizontalHeaderItem(i).text() for i in range(6)]
        self.assertEqual(actual_cols, expected_cols)

        # Check all materials are displayed in vertical scrolling (1,734 rows)
        self.assertEqual(view.table_catalog.rowCount(), 1734)
        self.assertEqual(view.table_catalog.item(0, 0).text(), "1")
        self.assertEqual(view.table_catalog.item(1, 0).text(), "2")
        row0_code = view.table_catalog.item(0, 1).text()
        self.assertEqual(row0_code, "1001NM210010BL421")

        # Test filtering
        view.filter_catalog("HOSE PIPE")
        self.assertGreater(len(view.filtered_catalog_items), 0)
        self.assertLess(len(view.filtered_catalog_items), 1734)
        # Check filtered S no starts from 1
        self.assertEqual(view.table_catalog.item(0, 0).text(), "1")

        # Test loading item into studio
        first_filtered = view.filtered_catalog_items[0]
        view.load_item_into_studio(first_filtered)
        self.assertEqual(view.stacked_widget.currentIndex(), 0)
        self.assertEqual(view.txt_mat_code.text(), first_filtered["material_code"])
        self.assertEqual(view.txt_mat_name.text(), first_filtered["material_name"])
        self.assertEqual(view.txt_uom.text(), first_filtered["uom"])

    def test_07_sticker_print_dialog_initialization(self):
        """Verify StickerPrintDialog initializes correctly with items."""
        sample_items = self.catalog[:2]
        dlg = StickerPrintDialog(sample_items)
        self.assertEqual(len(dlg.items), 2)
        self.assertIsNotNone(dlg.current_label_img)
        self.assertEqual(dlg.get_selected_size_code(), "50x25")


if __name__ == "__main__":
    unittest.main()
